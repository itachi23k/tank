import csv
from io import BytesIO, StringIO
from datetime import datetime, timedelta
from collections import defaultdict
from flask import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExcelGenerator:

    @staticmethod
    def exportar_equipamento_excel(equip, resultados, total_litros, data_inicio, data_fim):
        wb = Workbook()
        titulo_font = Font(name='Calibri', size=16, bold=True, color='1F4E79')
        cabecalho_font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
        cabecalho_fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
        dados_font = Font(name='Calibri', size=10, color='333333')
        total_fill = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')
        borda_fina = Border(
            left=Side(style='thin', color='B4C6E7'), right=Side(style='thin', color='B4C6E7'),
            top=Side(style='thin', color='B4C6E7'), bottom=Side(style='thin', color='B4C6E7')
        )

        ws = wb.active
        ws.title = "Extrato Detalhado"
        ws.merge_cells('A1:H1')
        ws['A1'] = 'DOSSIÊ ANALÍTICO DE ABASTECIMENTO'
        ws['A1'].font = titulo_font
        ws['A1'].alignment = Alignment(horizontal='center')
        ws.row_dimensions[1].height = 30

        ws.merge_cells('A3:H3')
        ws['A3'] = f'Equipamento: {equip.prefixo_placa} | Tipo: {equip.tipo_equipamento or "N/I"} | Locador: {equip.locador or "CENTRAL DE EQUIPAMENTOS"}'
        ws['A3'].font = Font(name='Calibri', size=11, bold=True, color='2E75B6')

        ws.merge_cells('A4:H4')
        ws['A4'] = f'Período: {datetime.strptime(data_inicio, "%Y-%m-%d").strftime("%d/%m/%Y")} a {datetime.strptime(data_fim, "%Y-%m-%d").strftime("%d/%m/%Y")}'

        ws.merge_cells('A5:H5')
        ws['A5'] = f'Total Geral: {total_litros:,.2f} Litros'
        ws['A5'].font = Font(name='Calibri', size=11, bold=True, color='C00000')

        linha = 7
        cabecalhos = ['Data', 'Dia', 'Fornecedor/Posto', 'Insumo', 'Volume (L)', 'Hodômetro', 'Horímetro', 'Observações']
        for col, cab in enumerate(cabecalhos, 1):
            cel = ws.cell(row=linha, column=col, value=cab)
            cel.font = cabecalho_font
            cel.fill = cabecalho_fill
            cel.alignment = Alignment(horizontal='center', vertical='center')
            cel.border = borda_fina

        dias_semana = ['Seg', 'Ter', 'Qua', 'Qui', 'Sex', 'Sáb', 'Dom']
        for i, item in enumerate(resultados):
            l = 8 + i
            ws.cell(row=l, column=1, value=item.data_abastecimento).number_format = 'DD/MM/YYYY'
            ws.cell(row=l, column=2, value=dias_semana[item.data_abastecimento.weekday()])
            ws.cell(row=l, column=3, value=item.fornecedor_rel.nome if item.fornecedor_rel else 'N/A')
            ws.cell(row=l, column=4, value=item.categoria_insumo)
            ws.cell(row=l, column=5, value=item.quantidade).number_format = '#,##0.00'
            ws.cell(row=l, column=6, value=item.hodometro or '')
            ws.cell(row=l, column=7, value=item.horimetro or '')
            ws.cell(row=l, column=8, value=item.descricao_obs or '')

            for col in range(1, 9):
                ws.cell(row=l, column=col).font = dados_font
                ws.cell(row=l, column=col).border = borda_fina

        linha_total = 8 + len(resultados)
        ws.merge_cells(f'A{linha_total}:D{linha_total}')
        ws.cell(row=linha_total, column=1, value='TOTAL').font = Font(name='Calibri', size=11, bold=True, color='1F4E79')
        ws.cell(row=linha_total, column=5, value=total_litros).font = Font(name='Calibri', size=12, bold=True, color='C00000')
        ws.cell(row=linha_total, column=5).number_format = '#,##0.00'
        for col in range(1, 9):
            ws.cell(row=linha_total, column=col).fill = total_fill
            ws.cell(row=linha_total, column=col).border = borda_fina

        for i, w in enumerate([14, 8, 30, 22, 14, 15, 15, 30], 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        ws.freeze_panes = 'A8'

        # Aba 2: Resumo
        ws2 = wb.create_sheet("Resumo por Insumo")
        resumo = defaultdict(float)
        for r in resultados:
            resumo[r.categoria_insumo] += r.quantidade

        ws2.merge_cells('A1:C1')
        ws2['A1'] = 'RESUMO POR TIPO DE INSUMO'
        ws2['A1'].font = titulo_font

        for col, cab in enumerate(['Insumo', 'Volume (L)', '%'], 1):
            cel = ws2.cell(row=3, column=col, value=cab)
            cel.font = cabecalho_font
            cel.fill = cabecalho_fill
            cel.border = borda_fina

        for i, (insumo, vol) in enumerate(sorted(resumo.items())):
            l = 4 + i
            ws2.cell(row=l, column=1, value=insumo).border = borda_fina
            ws2.cell(row=l, column=2, value=vol).number_format = '#,##0.00'
            ws2.cell(row=l, column=2).border = borda_fina
            ws2.cell(row=l, column=3, value=vol/total_litros if total_litros > 0 else 0).number_format = '0.0%'
            ws2.cell(row=l, column=3).border = borda_fina

        ws2.column_dimensions['A'].width = 25
        ws2.column_dimensions['B'].width = 18
        ws2.column_dimensions['C'].width = 12

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        nome = f'Dossie_{equip.prefixo_placa}_{data_inicio}_{data_fim}.xlsx'
        return Response(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': f'attachment; filename={nome}'}
        )

    @staticmethod
    def exportar_equipamento_csv(equip, resultados, total_litros, data_inicio, data_fim):
        si = StringIO()
        cw = csv.writer(si, delimiter=';')
        cw.writerow(['Data', 'Fornecedor', 'Insumo', 'Volume (L)', 'Hodometro', 'Horimetro', 'Observacoes'])
        for r in resultados:
            cw.writerow([
                r.data_abastecimento.strftime('%d/%m/%Y'),
                r.fornecedor_rel.nome if r.fornecedor_rel else 'N/A',
                r.categoria_insumo,
                str(r.quantidade).replace('.', ','),
                r.hodometro if r.hodometro else '',
                r.horimetro if r.horimetro else '',
                r.descricao_obs or ''
            ])
        cw.writerow([])
        cw.writerow(['TOTAL', '', '', str(total_litros).replace('.', ','), '', '', ''])
        output = si.getvalue()
        nome = f'Dossie_{equip.prefixo_placa}_{data_inicio}_{data_fim}.csv'
        return Response(
            '\ufeff' + output,
            mimetype='text/csv',
            headers={'Content-Disposition': f'attachment; filename={nome}'}
        )

    @staticmethod
    def exportar_equipamento_texto(equip, resultados, data_inicio, data_fim, insumo_filtro=None):
        dt_inicio = datetime.strptime(data_inicio, "%Y-%m-%d").date()
        dt_fim = datetime.strptime(data_fim, "%Y-%m-%d").date()

        if insumo_filtro and insumo_filtro != 'todos':
            resultados = [r for r in resultados if r.categoria_insumo.upper() == insumo_filtro.upper()]

        volumes_diarios = {}
        for r in resultados:
            volumes_diarios.setdefault(r.data_abastecimento, 0.0)
            volumes_diarios[r.data_abastecimento] += r.quantidade

        linhas = [
            f'# Valores Diários - {equip.prefixo_placa}',
            f'# Combustível: {insumo_filtro if insumo_filtro else "Todos"}',
            f'# Período: {dt_inicio.strftime("%d/%m/%Y")} a {dt_fim.strftime("%d/%m/%Y")}',
            ''
        ]
        data_atual = dt_inicio
        while data_atual <= dt_fim:
            if data_atual in volumes_diarios:
                linhas.append(f"{volumes_diarios[data_atual]:.2f}".replace('.', ','))
            else:
                linhas.append('')
            data_atual += timedelta(days=1)

        texto = '\n'.join(linhas)
        nome = f'Valores_{insumo_filtro or "Todos"}_{equip.prefixo_placa}.txt'
        return Response(texto, mimetype='text/plain; charset=utf-8', headers={'Content-Disposition': f'attachment; filename={nome}'})

    @staticmethod
    def exportar_fornecedor_excel(forn, abastecimentos, entradas_nf, transferencias, transferencias_enviadas, data_inicio, data_fim):
        wb = Workbook()
        titulo_font = Font(name='Calibri', size=13, bold=True, color='1F4E79')
        subtitulo_font = Font(name='Calibri', size=10, italic=True, color='555555')
        cabecalho_font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
        dados_font = Font(name='Calibri', size=10, color='333333')
        total_font = Font(name='Calibri', size=11, bold=True, color='1F4E79')
        total_fill = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')
        thin_border = Border(
            left=Side(style='thin', color='D0D9E4'), right=Side(style='thin', color='D0D9E4'),
            top=Side(style='thin', color='D0D9E4'), bottom=Side(style='thin', color='D0D9E4')
        )
        data_ini_format = datetime.strptime(data_inicio, '%Y-%m-%d').strftime('%d/%m/%Y')
        data_fim_format = datetime.strptime(data_fim, '%Y-%m-%d').strftime('%d/%m/%Y')

        # Aba 1: Saídas
        ws1 = wb.active
        ws1.title = "Saídas"
        fill_aba1 = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
        ws1['A1'] = f"EXTRATO DE SAÍDAS (ABASTECIMENTOS E TRANSBORDOS) - {forn.nome}"
        ws1['A1'].font = titulo_font
        ws1['A2'] = f"Período: {data_ini_format} a {data_fim_format}"
        ws1['A2'].font = subtitulo_font

        headers1 = ['Data', 'Tipo Saída', 'Destino / Placa', 'Descrição / Modelo', 'Insumo', 'Volume (L/Kg)', 'Hodômetro', 'Horímetro', 'Observações']
        ws1.append([])
        ws1.append(headers1)
        for col_idx in range(1, len(headers1) + 1):
            c = ws1.cell(row=4, column=col_idx)
            c.font = cabecalho_font
            c.fill = fill_aba1
            c.border = thin_border

        row_idx = 5
        tot_saidas = 0.0
        for a in abastecimentos:
            ws1.cell(row=row_idx, column=1, value=a.data_abastecimento.strftime('%d/%m/%Y')).alignment = Alignment(horizontal='center')
            ws1.cell(row=row_idx, column=2, value="Abastecimento").alignment = Alignment(horizontal='center')
            ws1.cell(row=row_idx, column=3, value=a.equipamento_rel.prefixo_placa if a.equipamento_rel else '-').alignment = Alignment(horizontal='center')
            ws1.cell(row=row_idx, column=4, value=a.equipamento_rel.descricao if a.equipamento_rel else '-')
            ws1.cell(row=row_idx, column=5, value=a.categoria_insumo)
            c_vol = ws1.cell(row=row_idx, column=6, value=round(a.quantidade, 2))
            c_vol.number_format = '#,##0.00'
            c_vol.alignment = Alignment(horizontal='right')
            ws1.cell(row=row_idx, column=7, value=a.hodometro or '').alignment = Alignment(horizontal='center')
            ws1.cell(row=row_idx, column=8, value=a.horimetro or '').alignment = Alignment(horizontal='center')
            ws1.cell(row=row_idx, column=9, value=a.descricao_obs or '-')
            for c in range(1, len(headers1) + 1):
                ws1.cell(row=row_idx, column=c).font = dados_font
                ws1.cell(row=row_idx, column=c).border = thin_border
            tot_saidas += a.quantidade
            row_idx += 1

        for te in transferencias_enviadas:
            dest_nome = te.destino.nome if te.destino else 'Outro Tanque'
            ws1.cell(row=row_idx, column=1, value=te.data_hora.strftime('%d/%m/%Y %H:%M')).alignment = Alignment(horizontal='center')
            ws1.cell(row=row_idx, column=2, value="Transbordo Emitido").alignment = Alignment(horizontal='center')
            ws1.cell(row=row_idx, column=3, value=dest_nome).alignment = Alignment(horizontal='center')
            ws1.cell(row=row_idx, column=4, value=f"Transbordo para {dest_nome}")
            ws1.cell(row=row_idx, column=5, value=te.tipo_combustivel)
            c_vol = ws1.cell(row=row_idx, column=6, value=round(te.litros, 2))
            c_vol.number_format = '#,##0.00'
            c_vol.alignment = Alignment(horizontal='right')
            ws1.cell(row=row_idx, column=7, value='-').alignment = Alignment(horizontal='center')
            ws1.cell(row=row_idx, column=8, value='-').alignment = Alignment(horizontal='center')
            ws1.cell(row=row_idx, column=9, value=te.observacao or '-')
            for c in range(1, len(headers1) + 1):
                ws1.cell(row=row_idx, column=c).font = dados_font
                ws1.cell(row=row_idx, column=c).border = thin_border
            tot_saidas += te.litros
            row_idx += 1

        ws1.cell(row=row_idx, column=1, value="TOTAL GERAL DE SAÍDAS")
        c_tot1 = ws1.cell(row=row_idx, column=6, value=round(tot_saidas, 2))
        c_tot1.number_format = '#,##0.00'
        for c in range(1, len(headers1) + 1):
            ws1.cell(row=row_idx, column=c).font = total_font
            ws1.cell(row=row_idx, column=c).fill = total_fill
            ws1.cell(row=row_idx, column=c).border = thin_border

        # Aba 2: Entradas NFs
        ws2 = wb.create_sheet(title="Entradas NFs")
        ws2['A1'] = f"ENTRADAS POR NOTA FISCAL - {forn.nome}"
        ws2['A1'].font = titulo_font
        headers2 = ['Data', 'Nº Nota Fiscal', 'Insumo', 'Quantidade (L/Kg)', 'Responsável']
        ws2.append([])
        ws2.append(headers2)
        row_idx2 = 5
        tot_nfs = 0.0
        for e in entradas_nf:
            ws2.cell(row=row_idx2, column=1, value=e.data_entrada.strftime('%d/%m/%Y'))
            ws2.cell(row=row_idx2, column=2, value=e.numero_nf or 'S/N')
            ws2.cell(row=row_idx2, column=3, value=e.categoria_insumo)
            c_v = ws2.cell(row=row_idx2, column=4, value=round(e.quantidade, 2))
            c_v.number_format = '#,##0.00'
            ws2.cell(row=row_idx2, column=5, value=e.usuario_rel.nome if e.usuario_rel else '-')
            tot_nfs += e.quantidade
            row_idx2 += 1

        # Aba 3: Transbordos Recebidos
        ws3 = wb.create_sheet(title="Entradas por Transbordo")
        ws3['A1'] = f"TRANSBORDOS RECEBIDOS - {forn.nome}"
        ws3['A1'].font = titulo_font
        headers3 = ['Data', 'Hora', 'Origem', 'Combustível', 'Volume (L)', 'Observação']
        ws3.append([])
        ws3.append(headers3)
        row_idx3 = 5
        for t in transferencias:
            ws3.cell(row=row_idx3, column=1, value=t.data_hora.strftime('%d/%m/%Y'))
            ws3.cell(row=row_idx3, column=2, value=t.data_hora.strftime('%H:%M'))
            ws3.cell(row=row_idx3, column=3, value=t.origem.nome if t.origem else '-')
            ws3.cell(row=row_idx3, column=4, value=t.tipo_combustivel)
            c_v = ws3.cell(row=row_idx3, column=5, value=round(t.litros, 2))
            c_v.number_format = '#,##0.00'
            ws3.cell(row=row_idx3, column=6, value=t.observacao or '-')
            row_idx3 += 1

        # Aba 4: Consolidado Geral
        ws4 = wb.create_sheet(title="Consolidado")
        ws4['A1'] = f"EXTRATO CONSOLIDADO - {forn.nome}"
        ws4['A1'].font = titulo_font
        headers4 = ['Data', 'Equipamento/Descrição', 'Entrada', 'Saída']
        ws4.append([])
        ws4.append(headers4)
        
        eventos = []
        for e in entradas_nf:
            eventos.append({'data': e.data_entrada, 'desc': f"Entrada NF: {e.numero_nf or 'S/N'}", 'ent': e.quantidade, 'sai': 0.0})
        for t in transferencias:
            eventos.append({'data': t.data_hora.date(), 'desc': f"Transbordo Recebido de {t.origem.nome if t.origem else '-'}", 'ent': t.litros, 'sai': 0.0})
        for te in transferencias_enviadas:
            eventos.append({'data': te.data_hora.date(), 'desc': f"Transbordo Emitido para {te.destino.nome if te.destino else '-'}", 'ent': 0.0, 'sai': te.litros})
        for a in abastecimentos:
            eventos.append({'data': a.data_abastecimento, 'desc': a.equipamento_rel.prefixo_placa if a.equipamento_rel else 'OUTROS', 'ent': 0.0, 'sai': a.quantidade})
        eventos.sort(key=lambda x: x['data'])

        row_idx4 = 5
        tot_ent = 0.0
        tot_sai = 0.0
        for ev in eventos:
            ws4.cell(row=row_idx4, column=1, value=ev['data'].strftime('%d/%m/%Y'))
            ws4.cell(row=row_idx4, column=2, value=ev['desc'])
            c_e = ws4.cell(row=row_idx4, column=3, value=round(ev['ent'], 2) if ev['ent'] > 0 else '-')
            c_s = ws4.cell(row=row_idx4, column=4, value=round(ev['sai'], 2) if ev['sai'] > 0 else '-')
            if ev['ent'] > 0: c_e.number_format = '#,##0.00'
            if ev['sai'] > 0: c_s.number_format = '#,##0.00'
            tot_ent += ev['ent']
            tot_sai += ev['sai']
            row_idx4 += 1

        for ws in [ws1, ws2, ws3, ws4]:
            for col in ws.columns:
                max_l = max(len(str(cell.value or '')) for cell in col)
                ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_l + 3, 12)

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        nome_arquivo = f'Extrato_Posto_{forn.nome}_{data_inicio}_{data_fim}.xlsx'.replace(' ', '_')
        return Response(buffer.getvalue(), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': f'attachment; filename={nome_arquivo}'})

    @staticmethod
    def exportar_consolidado_csv(forn, resultados, data_inicio, data_fim):
        si = StringIO()
        cw = csv.writer(si, delimiter=';')
        cw.writerow(['Posto/Distribuidor', 'Placa/Prefixo', 'Tipo de Equipamento', 'Volume Consolidado (L/Kg)'])
        for r in resultados:
            cw.writerow([
                forn.nome,
                r.prefixo_placa,
                r.tipo_equipamento,
                str(round(r.total_consumido, 2)).replace('.', ',')
            ])
        output = si.getvalue()
        nome_arquivo = f'Consolidado_{forn.nome}_{data_inicio}_{data_fim}.csv'.replace(' ', '_')
        return Response('\ufeff' + output, mimetype='text/csv', headers={'Content-Disposition': f'attachment; filename={nome_arquivo}'})

    @staticmethod
    def exportar_locador_excel(locador_sel, resultados_resumo, lancamentos_detalhados, total_litros, data_inicio, data_fim):
        wb = Workbook()
        titulo_font = Font(name='Calibri', size=14, bold=True, color='1F4E79')
        subtitulo_font = Font(name='Calibri', size=10, italic=True, color='555555')
        cabecalho_font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
        cabecalho_fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
        equip_header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
        equip_header_font = Font(name='Calibri', size=10, bold=True, color='1F4E79')
        subtotal_fill = PatternFill(start_color='EDEDED', end_color='EDEDED', fill_type='solid')
        subtotal_font = Font(name='Calibri', size=10, bold=True, color='1F4E79')
        dados_font = Font(name='Calibri', size=10, color='333333')
        total_fill = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')
        total_font = Font(name='Calibri', size=11, bold=True, color='1F4E79')
        thin_border = Border(
            left=Side(style='thin', color='D0D9E4'), right=Side(style='thin', color='D0D9E4'),
            top=Side(style='thin', color='D0D9E4'), bottom=Side(style='thin', color='D0D9E4')
        )
        data_ini_format = datetime.strptime(data_inicio, '%Y-%m-%d').strftime('%d/%m/%Y')
        data_fim_format = datetime.strptime(data_fim, '%Y-%m-%d').strftime('%d/%m/%Y')

        ws_det = wb.active
        ws_det.title = "Lançamentos por Equipamento"
        ws_det['A1'] = "RELATÓRIO DETALHADO DE CONSUMO POR LOCADOR"
        ws_det['A1'].font = titulo_font
        ws_det['A2'] = f"Locador: {locador_sel} | Período: {data_ini_format} a {data_fim_format}"
        ws_det['A2'].font = subtitulo_font

        headers_det = ['Data', 'Placa / Prefixo', 'Descrição / Modelo', 'Insumo', 'Posto / Origem', 'Volume (L)', 'Hodômetro', 'Horímetro', 'Observações']
        ws_det.append([])
        ws_det.append(headers_det)

        for col_idx in range(1, len(headers_det) + 1):
            c = ws_det.cell(row=4, column=col_idx)
            c.font = cabecalho_font
            c.fill = cabecalho_fill
            c.border = thin_border

        lancamentos_por_equip = defaultdict(list)
        for abast in lancamentos_detalhados:
            placa = abast.equipamento_rel.prefixo_placa if abast.equipamento_rel else 'OUTROS'
            lancamentos_por_equip[placa].append(abast)

        current_row = 5
        for eq_res in resultados_resumo:
            placa = eq_res.prefixo_placa
            abasts = lancamentos_por_equip.get(placa, [])
            ws_det.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers_det))
            cell_eq_title = ws_det.cell(row=current_row, column=1, value=f"🚜 EQUIPAMENTO: {placa} | Modelo: {eq_res.descricao or '-'} | Tipo: {eq_res.tipo_equipamento or '-'}")
            cell_eq_title.font = equip_header_font
            cell_eq_title.fill = equip_header_fill
            for c_idx in range(1, len(headers_det) + 1):
                ws_det.cell(row=current_row, column=c_idx).border = thin_border
            current_row += 1

            subtotal_eq = 0.0
            for ab in abasts:
                ws_det.cell(row=current_row, column=1, value=ab.data_abastecimento.strftime('%d/%m/%Y')).alignment = Alignment(horizontal='center')
                ws_det.cell(row=current_row, column=2, value=placa).alignment = Alignment(horizontal='center')
                ws_det.cell(row=current_row, column=3, value=eq_res.descricao or '-')
                ws_det.cell(row=current_row, column=4, value=ab.categoria_insumo)
                ws_det.cell(row=current_row, column=5, value=ab.fornecedor_rel.nome if ab.fornecedor_rel else '-')
                c_vol = ws_det.cell(row=current_row, column=6, value=round(ab.quantidade, 2))
                c_vol.number_format = '#,##0.00'
                c_vol.alignment = Alignment(horizontal='right')
                ws_det.cell(row=current_row, column=7, value=ab.hodometro or '').alignment = Alignment(horizontal='center')
                ws_det.cell(row=current_row, column=8, value=ab.horimetro or '').alignment = Alignment(horizontal='center')
                ws_det.cell(row=current_row, column=9, value=ab.descricao_obs or '-')

                for col_idx in range(1, len(headers_det) + 1):
                    c = ws_det.cell(row=current_row, column=col_idx)
                    c.font = dados_font
                    c.border = thin_border
                subtotal_eq += ab.quantidade
                current_row += 1

            ws_det.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
            ws_det.cell(row=current_row, column=1, value=f"TOTAL DO EQUIPAMENTO ({placa}) - {len(abasts)} abastecimentos").alignment = Alignment(horizontal='right', vertical='center')
            c_subtot = ws_det.cell(row=current_row, column=6, value=round(subtotal_eq, 2))
            c_subtot.number_format = '#,##0.00'
            c_subtot.alignment = Alignment(horizontal='right')
            for col_idx in range(1, len(headers_det) + 1):
                c = ws_det.cell(row=current_row, column=col_idx)
                c.font = subtotal_font
                c.fill = subtotal_fill
                c.border = thin_border
            current_row += 1

        ws_det.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
        ws_det.cell(row=current_row, column=1, value=f"TOTAL GERAL ({locador_sel})").alignment = Alignment(horizontal='right', vertical='center')
        c_tot_det = ws_det.cell(row=current_row, column=6, value=round(total_litros, 2))
        c_tot_det.number_format = '#,##0.00'
        for col_idx in range(1, len(headers_det) + 1):
            c = ws_det.cell(row=current_row, column=col_idx)
            c.font = total_font
            c.fill = total_fill
            c.border = thin_border

        for col in ws_det.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws_det.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 3, 12)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        nome = f'Relatorio_Locador_{locador_sel}_{data_inicio}_{data_fim}.xlsx'.replace(' ', '_')
        return Response(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': f'attachment; filename={nome}'})

    @staticmethod
    def exportar_resumo_mensal_excel(dados):
        wb = Workbook()
        
        # Paleta de cores moderna
        title_font = Font(name='Calibri', size=15, bold=True, color='1E293B')
        subtitle_font = Font(name='Calibri', size=10, italic=True, color='64748B')
        section_font = Font(name='Calibri', size=11, bold=True, color='1E3A8A')
        header_font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
        dados_font = Font(name='Calibri', size=9, color='1E293B')
        total_font = Font(name='Calibri', size=10, bold=True, color='0F172A')

        fill_header_dark = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
        fill_header_blue = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
        fill_header_green = PatternFill(start_color='059669', end_color='059669', fill_type='solid')
        fill_header_amber = PatternFill(start_color='D97706', end_color='D97706', fill_type='solid')
        fill_total = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')

        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1')
        )

        obra_nome = dados['obra'].nome if dados['obra'] else 'Todas as Obras'
        obra_cod = dados['obra'].codigo if dados['obra'] else 'GLOBAL'
        nome_mes = dados['nome_mes']
        ano = dados['ano']
        mes = dados['mes']
        nome_mes_ant = dados['nome_mes_anterior']

        # ----------------------------------------------------
        # ABA 1: BALANÇO MENSAL (CONSOLIDADO E TANQUES)
        # ----------------------------------------------------
        ws1 = wb.active
        ws1.title = "Balanço Mensal"

        # Cabeçalho da Planilha
        ws1.merge_cells('A1:H1')
        ws1['A1'] = f"FECHAMENTO MENSAL DE COMBUSTÍVEIS - {nome_mes.upper()} / {ano}"
        ws1['A1'].font = title_font
        ws1['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws1.row_dimensions[1].height = 28

        ws1.merge_cells('A2:H2')
        ws1['A2'] = f"Obra: {obra_cod} - {obra_nome} | Período: {dados['data_inicio'].strftime('%d/%m/%Y')} a {dados['data_fim'].strftime('%d/%m/%Y')}"
        ws1['A2'].font = subtitle_font
        ws1['A2'].alignment = Alignment(horizontal='center', vertical='center')

        # Seção 1: Resumo Consolidado por Combustível
        ws1.cell(row=4, column=1, value="1. CONSOLIDAÇÃO POR TIPO DE COMBUSTÍVEL").font = section_font
        
        headers_comb = [
            'Tipo de Combustível',
            f'Saldo Inicial ({nome_mes_ant})',
            '(+) Entradas NF',
            '(+/-) Transb. Líquido',
            '(-) Saídas (Abast.)',
            '(=) Saldo Final',
            'Capacidade Total',
            '% Ocupação Final'
        ]

        for c_idx, h in enumerate(headers_comb, 1):
            cell = ws1.cell(row=5, column=c_idx, value=h)
            cell.font = header_font
            cell.fill = fill_header_blue
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = thin_border
        ws1.row_dimensions[5].height = 25

        row = 6
        for c in dados['resumo_combustiveis']:
            ws1.cell(row=row, column=1, value=c['tipo']).alignment = Alignment(horizontal='left')
            ws1.cell(row=row, column=2, value=c['saldo_anterior']).number_format = '#,##0.00'
            ws1.cell(row=row, column=3, value=c['entradas_nf']).number_format = '#,##0.00'
            ws1.cell(row=row, column=4, value=c['transf_liquida']).number_format = '#,##0.00'
            ws1.cell(row=row, column=5, value=c['saidas']).number_format = '#,##0.00'
            ws1.cell(row=row, column=6, value=c['saldo_final']).number_format = '#,##0.00'
            ws1.cell(row=row, column=7, value=c['capacidade']).number_format = '#,##0.00'
            ws1.cell(row=row, column=8, value=f"{c['percentual']:.1f}%").alignment = Alignment(horizontal='center')

            for col_idx in range(1, 9):
                cell = ws1.cell(row=row, column=col_idx)
                cell.font = dados_font
                cell.border = thin_border
                if col_idx in [2, 3, 4, 5, 6, 7]:
                    cell.alignment = Alignment(horizontal='right')
            row += 1

        # Linha de Totais da Consolidação
        ws1.cell(row=row, column=1, value="TOTAL GERAL").alignment = Alignment(horizontal='left')
        ws1.cell(row=row, column=2, value=dados['totais']['saldo_anterior']).number_format = '#,##0.00'
        ws1.cell(row=row, column=3, value=dados['totais']['entradas']).number_format = '#,##0.00'
        ws1.cell(row=row, column=4, value=dados['totais']['transf_rec'] - dados['totais']['transf_env']).number_format = '#,##0.00'
        ws1.cell(row=row, column=5, value=dados['totais']['saidas']).number_format = '#,##0.00'
        ws1.cell(row=row, column=6, value=dados['totais']['saldo_final']).number_format = '#,##0.00'
        ws1.cell(row=row, column=7, value=dados['totais']['capacidade_total']).number_format = '#,##0.00'
        perc_tot = (dados['totais']['saldo_final'] / dados['totais']['capacidade_total'] * 100) if dados['totais']['capacidade_total'] > 0 else 0
        ws1.cell(row=row, column=8, value=f"{perc_tot:.1f}%").alignment = Alignment(horizontal='center')

        for col_idx in range(1, 9):
            cell = ws1.cell(row=row, column=col_idx)
            cell.font = total_font
            cell.fill = fill_total
            cell.border = thin_border
            if col_idx in [2, 3, 4, 5, 6, 7]:
                cell.alignment = Alignment(horizontal='right')
        row += 3

        # Seção 2: Balanço Físico por Tanque / Comboio
        ws1.cell(row=row, column=1, value="2. BALANÇO FÍSICO POR TANQUE / COMBOIO").font = section_font
        row += 1

        headers_tanques = [
            'Tanque / Comboio',
            'Categoria',
            'Combustível',
            'Capacidade (L)',
            f'Saldo Inicial ({nome_mes_ant})',
            '(+) Entradas NF',
            '(+) Transb. Rec.',
            '(-) Saídas (Abast.)',
            '(-) Transb. Env.',
            '(=) Saldo Final',
            '% Nível'
        ]

        for c_idx, h in enumerate(headers_tanques, 1):
            cell = ws1.cell(row=row, column=c_idx, value=h)
            cell.font = header_font
            cell.fill = fill_header_dark
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = thin_border
        ws1.row_dimensions[row].height = 25
        row += 1

        for t in dados['balanco_tanques']:
            ws1.cell(row=row, column=1, value=t['nome']).alignment = Alignment(horizontal='left')
            ws1.cell(row=row, column=2, value=t['categoria']).alignment = Alignment(horizontal='center')
            ws1.cell(row=row, column=3, value=t['tipo_combustivel']).alignment = Alignment(horizontal='center')
            ws1.cell(row=row, column=4, value=t['capacidade']).number_format = '#,##0.00'
            ws1.cell(row=row, column=5, value=t['saldo_anterior']).number_format = '#,##0.00'
            ws1.cell(row=row, column=6, value=t['entradas_nf']).number_format = '#,##0.00'
            ws1.cell(row=row, column=7, value=t['transf_rec']).number_format = '#,##0.00'
            ws1.cell(row=row, column=8, value=t['saidas']).number_format = '#,##0.00'
            ws1.cell(row=row, column=9, value=t['transf_env']).number_format = '#,##0.00'
            ws1.cell(row=row, column=10, value=t['saldo_final']).number_format = '#,##0.00'
            ws1.cell(row=row, column=11, value=f"{t['percentual']:.1f}%").alignment = Alignment(horizontal='center')

            for col_idx in range(1, 12):
                cell = ws1.cell(row=row, column=col_idx)
                cell.font = dados_font
                cell.border = thin_border
                if col_idx in [4, 5, 6, 7, 8, 9, 10]:
                    cell.alignment = Alignment(horizontal='right')
            row += 1

        for col in ws1.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws1.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 3, 14)

        # ----------------------------------------------------
        # ABA 2: NOTAS FISCAIS RECEBIDAS NO MÊS
        # ----------------------------------------------------
        ws2 = wb.create_sheet("Notas Fiscais")
        ws2.merge_cells('A1:F1')
        ws2['A1'] = f"ENTRADAS DE NOTAS FISCAIS - {nome_mes.upper()} / {ano}"
        ws2['A1'].font = title_font
        ws2['A1'].alignment = Alignment(horizontal='center', vertical='center')

        headers_nfs = ['Data', 'Nº NF', 'Fornecedor / Destino', 'Combustível / Insumo', 'Volume (L/Kg)', 'Cadastrado Por']
        for c_idx, h in enumerate(headers_nfs, 1):
            cell = ws2.cell(row=3, column=c_idx, value=h)
            cell.font = header_font
            cell.fill = fill_header_green
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border

        row_nf = 4
        tot_nf_vol = 0.0
        for nf in dados['entradas_nfs']:
            ws2.cell(row=row_nf, column=1, value=nf.data_entrada.strftime('%d/%m/%Y')).alignment = Alignment(horizontal='center')
            ws2.cell(row=row_nf, column=2, value=nf.numero_nf or 'S/N').alignment = Alignment(horizontal='center')
            ws2.cell(row=row_nf, column=3, value=nf.fornecedor_rel.nome if nf.fornecedor_rel else '-').alignment = Alignment(horizontal='left')
            ws2.cell(row=row_nf, column=4, value=nf.categoria_insumo).alignment = Alignment(horizontal='left')
            c_v = ws2.cell(row=row_nf, column=5, value=nf.quantidade)
            c_v.number_format = '#,##0.00'
            c_v.alignment = Alignment(horizontal='right')
            ws2.cell(row=row_nf, column=6, value=nf.usuario_rel.nome if nf.usuario_rel else '-').alignment = Alignment(horizontal='left')

            for col_idx in range(1, 7):
                ws2.cell(row=row_nf, column=col_idx).font = dados_font
                ws2.cell(row=row_nf, column=col_idx).border = thin_border
            tot_nf_vol += nf.quantidade
            row_nf += 1

        ws2.merge_cells(start_row=row_nf, start_column=1, end_row=row_nf, end_column=4)
        ws2.cell(row=row_nf, column=1, value=f"TOTAL DE NOTAS FISCAIS ({len(dados['entradas_nfs'])} NFs)").alignment = Alignment(horizontal='right')
        c_tot_nf = ws2.cell(row=row_nf, column=5, value=tot_nf_vol)
        c_tot_nf.number_format = '#,##0.00'
        c_tot_nf.alignment = Alignment(horizontal='right')

        for col_idx in range(1, 7):
            cell = ws2.cell(row=row_nf, column=col_idx)
            cell.font = total_font
            cell.fill = fill_total
            cell.border = thin_border

        for col in ws2.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws2.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 3, 14)

        # ----------------------------------------------------
        # ABA 3: CONSUMO POR EQUIPAMENTO NO MÊS
        # ----------------------------------------------------
        ws3 = wb.create_sheet("Consumo por Equipamento")
        ws3.merge_cells('A1:G1')
        ws3['A1'] = f"CONSUMO POR EQUIPAMENTO - {nome_mes.upper()} / {ano}"
        ws3['A1'].font = title_font
        ws3['A1'].alignment = Alignment(horizontal='center', vertical='center')

        headers_eq = ['Placa/Prefixo', 'Descrição / Modelo', 'Tipo Equipamento', 'Locador', 'Insumo', 'Qtd Abast.', 'Total Consumido (L)']
        for c_idx, h in enumerate(headers_eq, 1):
            cell = ws3.cell(row=3, column=c_idx, value=h)
            cell.font = header_font
            cell.fill = fill_header_amber
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border

        row_eq = 4
        tot_eq_vol = 0.0
        tot_eq_abast = 0
        for eq in dados['consumo_equipamentos']:
            ws3.cell(row=row_eq, column=1, value=eq.prefixo_placa).alignment = Alignment(horizontal='center')
            ws3.cell(row=row_eq, column=2, value=eq.descricao or '-').alignment = Alignment(horizontal='left')
            ws3.cell(row=row_eq, column=3, value=eq.tipo_equipamento or '-').alignment = Alignment(horizontal='left')
            ws3.cell(row=row_eq, column=4, value=eq.locador or '-').alignment = Alignment(horizontal='left')
            ws3.cell(row=row_eq, column=5, value=eq.categoria_insumo).alignment = Alignment(horizontal='center')
            ws3.cell(row=row_eq, column=6, value=eq.total_abastecimentos).alignment = Alignment(horizontal='center')
            c_v = ws3.cell(row=row_eq, column=7, value=eq.total_litros)
            c_v.number_format = '#,##0.00'
            c_v.alignment = Alignment(horizontal='right')

            for col_idx in range(1, 8):
                ws3.cell(row=row_eq, column=col_idx).font = dados_font
                ws3.cell(row=row_eq, column=col_idx).border = thin_border
            tot_eq_vol += eq.total_litros
            tot_eq_abast += eq.total_abastecimentos
            row_eq += 1

        ws3.merge_cells(start_row=row_eq, start_column=1, end_row=row_eq, end_column=5)
        ws3.cell(row=row_eq, column=1, value="TOTAL GERAL").alignment = Alignment(horizontal='right')
        ws3.cell(row=row_eq, column=6, value=tot_eq_abast).alignment = Alignment(horizontal='center')
        c_tot_eq = ws3.cell(row=row_eq, column=7, value=tot_eq_vol)
        c_tot_eq.number_format = '#,##0.00'
        c_tot_eq.alignment = Alignment(horizontal='right')

        for col_idx in range(1, 8):
            cell = ws3.cell(row=row_eq, column=col_idx)
            cell.font = total_font
            cell.fill = fill_total
            cell.border = thin_border

        for col in ws3.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws3.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 3, 14)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        nome_arq = f"Fechamento_Mensal_{obra_cod}_{ano}_{mes:02d}.xlsx".replace(' ', '_')
        return Response(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': f'attachment; filename={nome_arq}'})

